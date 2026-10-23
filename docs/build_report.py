"""Build the report from saved experimental artifacts without training models.

Requires python-docx and Pillow. Run from any directory with Python 3.12+.
DOCX PDF export is separate; see docs/report_generation.md.
"""
from __future__ import annotations

import csv
import hashlib
import json
import shutil
from pathlib import Path

from docx import Document
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor

ROOT = Path(__file__).resolve().parents[1]
OUT = ROOT / "deliverables/report"
DOCS = ROOT / "docs"
TITLE = "LSTM Based Internet Traffic Forecasting for Intelligent Bandwidth Prediction in Cloud ISP and Data Center Networks"
MODELS = ["persistence", "moving_average_4", "lstm", "gru", "transformer"]
LABELS = {"persistence": "Persistence", "moving_average_4": "Moving average", "lstm": "LSTM", "gru": "GRU", "transformer": "Transformer"}
sources: dict[str, str] = {}
md: list[str] = []


def source(path: Path) -> Path:
    sources[path.relative_to(ROOT).as_posix()] = hashlib.sha256(path.read_bytes()).hexdigest()
    return path


def rows(path: Path) -> list[dict[str, str]]:
    with source(path).open(encoding="utf-8-sig", newline="") as handle:
        return list(csv.DictReader(handle))


def config(path: Path) -> dict:
    return json.loads(source(path).read_text(encoding="utf-8"))


def para(text: str, style: str | None = None):
    p = doc.add_paragraph(text, style)
    md.append(text + "\n")
    return p


def heading(text: str, level: int = 1):
    doc.add_heading(text, level)
    md.append("#" * (level + 1) + " " + text + "\n")


def page(title: str):
    doc.add_page_break()
    heading(title)


def table(headers: list[str], data: list[list[str]], widths: list[float] | None = None):
    t = doc.add_table(rows=1, cols=len(headers))
    t.alignment = WD_TABLE_ALIGNMENT.CENTER
    t.autofit = False
    if widths is None:
        widths = [6.85 / len(headers)] * len(headers)
    for col, w in zip(t.columns, widths):
        col.width = Inches(w)
    for c, title in zip(t.rows[0].cells, headers):
        c.text = title
    for row in data:
        for c, value in zip(t.add_row().cells, row):
            c.text = str(value)
    for i, row in enumerate(t.rows):
        trPr = row._tr.get_or_add_trPr()
        cant = OxmlElement("w:cantSplit")
        trPr.append(cant)
        if i == 0:
            repeat = OxmlElement("w:tblHeader")
            trPr.append(repeat)
        for j, cell in enumerate(row.cells):
            cell.width = Inches(widths[j])
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            shading = OxmlElement("w:shd")
            shading.set(qn("w:fill"), "17324D" if i == 0 else ("F1F4F6" if i % 2 else "FFFFFF"))
            tcPr.append(shading)
            borders = OxmlElement("w:tcBorders")
            for edge in ("top", "left", "bottom", "right"):
                e = OxmlElement("w:" + edge)
                for k, v in {"val": "single", "sz": "4", "color": "D9D9D9"}.items():
                    e.set(qn("w:" + k), v)
                borders.append(e)
            tcPr.append(borders)
            margins = OxmlElement("w:tcMar")
            for edge, val in (("top", "65"), ("bottom", "65"), ("left", "90"), ("right", "90")):
                e = OxmlElement("w:" + edge)
                e.set(qn("w:w"), val)
                e.set(qn("w:type"), "dxa")
                margins.append(e)
            tcPr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(0)
                p.paragraph_format.space_before = Pt(0)
                p.paragraph_format.line_spacing = 1.05
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT if j == 0 else WD_ALIGN_PARAGRAPH.CENTER
                for run in p.runs:
                    run.font.size = Pt(9)
                    run.bold = i == 0
                    run.font.color.rgb = RGBColor.from_string("FFFFFF" if i == 0 else "000000")
    md.extend(["| " + " | ".join(headers) + " |", "| " + " | ".join(["---"] * len(headers)) + " |"])
    md.extend("| " + " | ".join(map(str, row)) + " |" for row in data)
    md.append("")
    doc.add_paragraph().paragraph_format.space_after = Pt(1)


def image(path: Path, caption: str, width=6.8):
    source(path)
    local = OUT / path.name
    if path.resolve() != local.resolve():
        shutil.copyfile(path, local)
    p = doc.add_paragraph()
    p.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p.paragraph_format.space_after = Pt(2)
    p.add_run().add_picture(str(local), width=Inches(width))
    cap = doc.add_paragraph(caption, "Caption")
    cap.paragraph_format.space_after = Pt(8)
    md.extend([f"![{caption}](../deliverables/report/{local.name})", caption + "\n"])


def code(lines: list[str]):
    p = doc.add_paragraph()
    p.paragraph_format.space_after = Pt(8)
    for i, line in enumerate(lines):
        r = p.add_run(("\n" if i else "") + line)
        r.font.name = "Consolas"
        r.font.size = Pt(8.5)
    md.extend(["```text", *lines, "```", ""])


def metric_data(split: str):
    return [[str(h), LABELS[m], f"{float(v['mae_mbps']):.3f}", f"{float(v['rmse_mbps']):.3f}", f"{float(v['mape_pct']):.2f}", f"{float(v['r2']):.3f}"] for h in (10, 30, 60) for m in MODELS for v in [metrics[h, m, split]]]


OUT.mkdir(parents=True, exist_ok=True)
DOCS.mkdir(parents=True, exist_ok=True)
metrics = {}
configs = {}
for h in (10, 30, 60):
    for family in ("baselines", "lstm", "gru", "transformer"):
        for row in rows(ROOT / f"results/{family}/{h}min/metrics.csv"):
            metrics[h, row["model"], row["split"]] = row
        configs[h, family] = config(ROOT / f"results/{family}/{h}min/run_config.json")
if len(metrics) != 30:
    raise ValueError("Expected five predictors across three horizons and two splits")
prep = config(ROOT / "data/prepared/preparation_config.json")
counts = rows(ROOT / "data/prepared/window_counts.csv")
allocation = rows(ROOT / "results/allocation/10min/metrics.csv")
defaults = rows(ROOT / "results/allocation/10min/default_test_metrics.csv")
ac = config(ROOT / "results/allocation/10min/run_config.json")
if len(allocation) != 40 or len(defaults) != 5:
    raise ValueError("Report requires the completed five-policy allocation comparison")
default_map = {r["policy"]: r for r in defaults}
policies = ["reactive", "proactive_moving_average_4", "proactive_lstm", "proactive_gru", "proactive_transformer"]
plabel = dict(zip(policies, ["Reactive", "Moving average", "LSTM", "GRU", "Transformer"]))
doc = Document()
sec = doc.sections[0]
sec.page_width = Inches(8.27)
sec.page_height = Inches(11.69)
sec.top_margin = Inches(.70)
sec.bottom_margin = Inches(.70)
sec.left_margin = Inches(.71)
sec.right_margin = Inches(.71)
sec.footer_distance = Inches(.3)
for name in ("Normal", "Title", "Subtitle", "Heading 1", "Heading 2", "Caption"):
    s = doc.styles[name]
    s.font.name = "Calibri"
    s.font.color.rgb = RGBColor(0, 0, 0)
    s.font.size = Pt(11)
    s.paragraph_format.space_after = Pt(7)
    s.paragraph_format.line_spacing = 1.12
doc.styles["Title"].font.size = Pt(25)
doc.styles["Title"].font.bold = True
doc.styles["Title"].paragraph_format.space_after = Pt(18)
doc.styles["Heading 1"].font.size = Pt(17)
doc.styles["Heading 1"].paragraph_format.space_after = Pt(12)
doc.styles["Heading 2"].font.size = Pt(12)
doc.styles["Heading 2"].font.bold = True
doc.styles["Caption"].font.size = Pt(9)
doc.styles["Caption"].font.italic = True
# The bundled Word template may attach a blue border to Title. Clear all
# inherited paragraph borders so the report title remains plain black text.
for style in doc.styles:
    for border in list(style.element.iter(qn("w:pBdr"))):
        border.getparent().remove(border)
footer = sec.footer.paragraphs[0]
footer.alignment = WD_ALIGN_PARAGRAPH.RIGHT
r = footer.add_run("Project report  |  ")
r.font.size = Pt(8)
field = OxmlElement("w:fldSimple")
field.set(qn("w:instr"), "PAGE")
footer._p.append(field)
doc.core_properties.title = TITLE
doc.core_properties.subject = "CESNET institution 0 forecasting and offline allocation experiment"
doc.core_properties.author = ""
doc.core_properties.keywords = "CESNET; LSTM; GRU; Transformer; forecasting; bandwidth allocation"

doc.add_paragraph(TITLE, "Title")
md.append("# " + TITLE + "\n")
para("Final academic project report", "Subtitle")
para("22 October 2026")
heading("Abstract")
para("This project implements a reproducible pipeline for forecasting recorded Internet traffic and evaluating proactive capacity allocation. The experiment uses institution 0 from CESNET-TimeSeries24 version v1, with ten-minute observations and three hours of historical inputs. Persistence and a four-observation moving average are compared with compact LSTM, GRU and Transformer models at direct forecast horizons of 10, 30 and 60 minutes. Chronological evaluation, training-only scaling and explicit exclusion of suspect measurement periods control avoidable leakage.")
para("On the held-out test data, Transformer achieves the lowest ten-minute RMSE at 30.803 Mbps. GRU achieves the lowest thirty-minute and sixty-minute RMSE at 32.290 and 34.785 Mbps. LSTM remains competitive but does not win every forecasting comparison. These are descriptive results for one institution and one fixed seed, without confidence intervals.")
lstm = default_map["proactive_lstm"]
reactive = default_map["reactive"]
para(f"In the ten-minute allocation simulation, LSTM with a 20% margin has shortfalls in {float(lstm['shortfall_pct']):.2f}% of test intervals, compared with {float(reactive['shortfall_pct']):.2f}% for reactive allocation. Its mean allocated capacity is {float(lstm['mean_capacity_mbps']):.2f} Mbps versus {float(reactive['mean_capacity_mbps']):.2f} Mbps for reactive allocation. Lower shortfall therefore comes with additional capacity. None of the five policies meets the target of at most 1% shortfall intervals on validation at the tested margins, so 20% is a disclosed fallback.")
heading("Experimental scope", 2)
para("The deliverable is an offline research prototype with saved forecasts, model checkpoints, result tables, a local dashboard and a reproducible workflow. The recorded traffic target is an institution-level nominal rate. Cloud and data-center deployment are proposed extensions; physical-link utilization, packet loss and latency improvements are not measured here.")

page("Problem statement and related research")
para("Bandwidth provisioning must trade unused capacity against demand that exceeds available capacity. A reactive rule uses the latest measurement and can lag a sudden change. A forecasting model can use recent temporal patterns to estimate upcoming demand, but an improved average forecast error does not automatically imply a reliable capacity policy. This study evaluates prediction and allocation separately so that both benefits and costs remain visible.")
heading("Research questions", 2)
para("The experiment asks whether compact neural models improve on simple predictors for the same future targets; how performance changes with forecast lead; and whether a forecast-based capacity rule reduces shortfall at identical margins. The study also examines whether the validation-selected policy can achieve a stated reliability requirement before reporting performance on held-out observations.")
heading("Model foundations", 2)
para("Hochreiter and Schmidhuber introduced LSTM to address difficulties learning information over long time lags in recurrent networks [3]. The implemented PyTorch LSTM is a modern gated recurrent layer, used here to encode an observed traffic window and produce one scalar forecast. This project does not reproduce the original paper's artificial-sequence experiments.")
para("Cho et al. proposed a recurrent encoder-decoder with gated hidden units in machine translation [4]. GRU provides a compact recurrent alternative in this study. Vaswani et al. introduced a Transformer based on attention [5]. Here a small encoder-only adaptation attends to historical traffic features; it is not the original translation architecture. These publications motivate model families, while all traffic accuracy claims in this report come from the local experiments.")
heading("Deliverable requirements", 2)
table(["Requirement", "Implementation"], [["Real traffic data", "CESNET v1 institution 0 and explicit provenance"], ["Comparable predictors", "Two baselines and three neural architectures"], ["Multiple leads", "Separate direct models for 10, 30 and 60 minutes"], ["Proactive resource study", "Ten-minute allocation at four equal margins"], ["Reviewable demonstration", "Saved artifacts, dashboard and historical replay"], ["Reproducibility", "Configuration, source hashes and validation commands"]], [2.0, 4.85])
para("The original project direction used GÉANT. The final implementation uses CESNET after a documented dataset change. GÉANT-specific units, timing and earlier draft performance claims are excluded from the experimental findings.")

page("Dataset provenance and measurement meaning")
para("CESNET-TimeSeries24 provides traffic time series collected from the CESNET3 network. The dataset record supplies institution, subnet and IP-level data and a ten-minute aggregation option [1]. The associated paper documents collection and limitations [2]. This experiment pins Zenodo record 13382427, version v1, rather than silently substituting a later dataset release.")
table(["Property", "Value or interpretation"], [["Selected series", "Institution 0 at ten-minute aggregation"], ["Raw observations", "40,298 rows with 19 columns"], ["First supplied timestamp", "2023-10-09 02:03:49 +02:00"], ["Last supplied timestamp", "2024-07-14 23:50:52 +02:00"], ["Join key", "id_time matched to times_10_minutes.csv"], ["Forecast target", "Recorded bidirectional n_bytes expressed as nominal Mbps"], ["Data license", "Creative Commons Attribution 4.0 International"]], [2.1, 4.75])
para("The project defines recorded Mbps as the observed byte count multiplied by eight and divided by 600,000,000. The 600-second denominator is the nominal interval length. An irregular timestamp gap is not used as a substitute measurement duration. The target is neither link utilization nor a directional download rate. The collection logic can update internal endpoints; institution exports have not been established to form mutually exclusive traffic totals. Consequently the study does not sum institutions into unique network-wide demand.")
heading("Selection and quality audit", 2)
para("Before fitting any models, the project compared all 283 institutions for byte-value quality, mapping coverage and prospective window availability. Fifty-six had all supplied time IDs and 229 had at least 95% mapped-row coverage. Institution 0 tied with several candidates on usable-window availability and had already received a complete sample audit. Its selection was based on availability, not forecasting accuracy.")
para("The institution sample has no duplicate or unmapped time IDs, and its measurement cells contain no blank, nonnumeric, nonfinite, negative or zero values. Nevertheless, five UTC timestamp transitions differ from 600 seconds, producing six continuous blocks. Populated cells and regular timestamps cannot establish complete capture. These local audit findings are preserved in the inspection outputs and research notes.")
para("Two archive members supply the experiment: institutions/agg_10_minutes/0.csv and times/times_10_minutes.csv. The first is renamed institution_0.csv without changing its contents. Source hashes in preparation_config.json identify the exact samples used. Attribution and the CC BY 4.0 license must accompany redistributed data extracts.")

page("Preprocessing and chronological windows")
para("The paper reports a monitoring failure around 21 May to 4 June 2024 [2]. Because exact institution impact and incident timezone are unresolved, the project excludes all supplied-calendar dates from 20 May through 5 June inclusive. This conservative policy marks 2,448 observations across 17 days. It is a project decision rather than an author-provided incident label. Raw values remain available for inspection.")
para("Both endpoints of each uninterrupted timestamp block are also guarded because boundary measurement durations are uncertain. These 12 observations do not overlap the outage dates, leaving 37,838 eligible observations. Every modeling path, from its earliest historical input through its future target, must remain in one regular block and avoid excluded observations. The pipeline never joins observations across an excluded interval.")
heading("Causal features and target availability", 2)
para("Each input contains 18 historical records, representing three hours. Its eight features are recorded demand; hour sine and cosine; weekday sine and cosine; an input-missing flag; and a trailing mean and population standard deviation over up to six observations within that input window. Calendar features use the supplied clock offset. Rolling calculations use only historical values through their current input position.")
para("Missing input rates may carry forward an earlier observation within an allowed block, with history reset at exclusions. Missing targets are never filled. Institution 0 has no missing byte observations, so no imputation occurs in the actual experiment. No invented latency, active-user count or packet arrival rate enters the models.")
para("An interval is assumed available at its supplied start plus 600 seconds. The last input completion defines the forecast origin, and target completion is 10, 30 or 60 minutes later. This is an offline availability assumption: actual export delay was not verified. Separate direct models predict each horizon, without recursively feeding predictions into later forecasts.")
table(["Lead", "Train", "Validation", "Test", "Total windows"], [[r["horizon_minutes"] + " min", r["train"], r["validation"], r["test"], r["total"]] for r in counts], [1.15, 1.35, 1.4, 1.3, 1.65])
para("Chronological 70/15/15 cut points are set on the full ordered timeline before exclusions. Target completion assigns a window to its split. Validation starts at 2024-04-21 22:40:52 UTC and test at 2024-06-02 22:10:52 UTC. The outage removes the beginning of the nominal test period. Earlier observed history may enter a later split, consistent with forecasting from information already available.")
para("Each horizon has its own feature and target standard scalers fitted exclusively on training windows. Constant feature standard deviations become one. Saved scaler values and source signatures make this choice auditable. Model-ready arrays use float32; original-unit rates are retained for evaluation.")

page("Forecasting models and training protocol")
table(["Predictor", "Definition", "Parameters"], [["Persistence", "Repeat the latest observed input demand", "0"], ["Moving average", "Mean of the last four observations", "0"], ["LSTM", "Two layers with 64 hidden units", "52,289"], ["GRU", "Two layers with 64 hidden units", "39,233"], ["Transformer", "Two layers, dimension 64, four attention heads", "67,585"]], [1.4, 4.35, 1.1])
para("The four-observation moving average covers 40 nominal minutes. Both baselines use the same forecasting origins and targets as the neural models and are necessary checks against unnecessary model complexity. They have no fitted training stage.")
para("LSTM and GRU read the eight features in chronological order and send the final hidden output through a scalar linear head. They are unidirectional and initialize recurrent state independently for every window. The Transformer first projects eight features to dimension 64 and adds fixed sinusoidal position encodings for 18 positions. Its two independently constructed encoder layers use four heads, a 128-unit feed-forward block, GELU, post-normalization and no dropout. The last position supplies the scalar head.")
para("Transformer attention can see every position in the supplied historical window because all those observations exist at the forecast origin. It receives no future target token. This design does not create look-ahead leakage. Mathematical attention and a disabled inference fast path preserve the intended training/evaluation implementation.")
heading("Shared optimization settings", 2)
table(["Setting", "Value"], [["Seed", "42 for the reported experiment"], ["Optimizer and learning rate", "Adam, 0.001"], ["Batch size and loss", "64, scaled-target mean squared error"], ["Epoch limit and early stopping", "Maximum 30; patience 5; minimum improvement 0"], ["Checkpoint selection", "Lowest validation MSE, restored before test scoring"], ["Regularization", "Dropout 0; no additional tuning study"], ["Training platform", "RTX 4050 Laptop GPU; CPU execution also supported"]], [2.45, 4.4])
para("Training windows are shuffled using a seeded generator, while validation and test remain chronological. PyTorch deterministic algorithms are enabled, cuDNN benchmarking is disabled, and CUBLAS_WORKSPACE_CONFIG is set to :4096:8. The seed supports repeatability within the recorded environment but cannot establish statistical robustness or bitwise identity across platforms.")
para("Each model and horizon writes its own checkpoint, training history, scalers, configuration, timestamped predictions and metric CSV. Test values are evaluated after validation checkpoint selection. Hyperparameters were fixed across horizons, so a disappointing result remains part of the comparison.")

page("Evaluation protocol and validation results")
para("All forecasting errors are computed in original Mbps after inverse target scaling. MAE is mean absolute error, while RMSE is the square root of mean squared error and places greater weight on large misses. MAPE is the mean absolute percentage error restricted to positive observed targets; excluded-target counts are saved. R squared compares squared error with the variability around the observed mean. Undefined values remain undefined rather than receiving arbitrary replacements.")
para("Every target in the reported validation and test sets is positive, so MAPE excludes zero observations. Forecasts are not clipped for accuracy scoring. No negative forecasts occur in these saved runs. Every predictor within a horizon uses identical target IDs and actual demand. Slightly different valid target sets at different horizons prevent a perfectly paired cross-horizon comparison.")
table(["Lead min", "Predictor", "MAE Mbps", "RMSE Mbps", "MAPE %", "R squared"], metric_data("validation"), [.65, 1.45, 1.15, 1.25, 1.1, 1.25])
para("Table 1. Validation results, with 4,029 targets at every horizon. Values are loaded from each model family's metrics.csv; baselines share their family file. The saved checkpoint is selected by validation loss within each model, not by test performance.", "Caption")
para("The lowest neural validation RMSE belongs to GRU at ten minutes, Transformer at thirty minutes and GRU at sixty minutes. Validation and test rankings can differ. Selecting a model or margin after inspecting the test set would turn that set into a development set; this report instead presents all comparisons transparently.")

page("Held out forecasting results")
table(["Lead min", "Predictor", "MAE Mbps", "RMSE Mbps", "MAPE %", "R squared"], metric_data("test"), [.65, 1.45, 1.15, 1.25, 1.1, 1.25])
para("Table 2. Test results. Sample counts are 5,554 at ten minutes, 5,548 at thirty minutes and 5,539 at sixty minutes. Rounded values are presented here; original precision remains in the generated CSVs.", "Caption")
para("At ten minutes, Transformer has the lowest MAE and RMSE, whereas the moving average has the lowest MAPE. All three neural predictors improve RMSE over persistence. LSTM has slightly higher MAE than the moving average despite a lower RMSE, illustrating that the preferred predictor depends on the objective.")
para("At thirty minutes, GRU has the lowest MAE, RMSE and MAPE among the compared predictors. Transformer is close in MAE and RMSE. This small gap is descriptive; one seed and one test period cannot establish a reliable population-level advantage.")
para("At sixty minutes, GRU has the lowest MAE and RMSE, and LSTM has the lowest MAPE. The Transformer RMSE is only slightly below the moving average while its MAE and MAPE are worse. Model complexity therefore does not guarantee better forecasts, particularly when the horizon and traffic distribution change.")
para("The project satisfies its forecasting objective by producing a consistent, explainable comparison. It does not require LSTM to outperform every baseline or competing architecture. These results should replace performance claims in any earlier draft that were not produced by the implemented experiment.")

page("Horizon effects and computational cost")
figure = ROOT / "results/comparison/forecast_test_10_30_60min.png"
if figure.exists():
    image(figure, "Figure 1. Test error by horizon for the same five predictor families. Each horizon has its own eligible target set.")
para("MAE and RMSE increase with lead for every predictor in these results. Percentage errors need not change monotonically because MAPE weights low-demand targets strongly, and cross-horizon target sets differ. The observed pattern motivates reporting several measures rather than selecting only the most favorable one.")
timing = []
for h in (10, 30, 60):
    for m in ("lstm", "gru", "transformer"):
        c = configs[h, m]
        v = metrics[h, m, "test"]
        timing.append([str(h), LABELS[m], str(c["epochs_completed"]), str(c["best_epoch"]), f"{float(v['training_time_s']):.2f}", f"{float(v['cpu_inference_ms']):.3f}"])
table(["Lead min", "Model", "Epochs", "Best epoch", "Train seconds", "CPU ms"], timing, [.65, 1.45, .85, 1.1, 1.45, 1.35])
para("Table 3. Training and single-window CPU inference timings from the saved runs. Baseline training time is zero; baseline inference timings are retained in forecast_metrics.csv alongside all neural timings.", "Caption")
para("Training time covers the epoch loop, validation and improved-checkpoint writes, excluding loading and final evaluation. Inference uses one CPU thread, 25 warm-up calls and 1,000 timed single-window calls on preloaded features; inverse scaling is included. Loading, feature construction, application overhead and collector delay are excluded. Different stopping epochs and machine load affect elapsed times, so these values are not a controlled hardware benchmark or an end-to-end latency guarantee.")

page("Proactive allocation experiment")
para("The allocation experiment uses ten-minute forecasts only, matching the earliest prediction horizon. Every policy provisions an abstract capacity pool for one 600-second interval. The reactive policy uses the latest observed demand, equivalent to a persistence forecast. Four proactive policies use moving-average, LSTM, GRU or Transformer forecasts. All policies receive exactly the same 5%, 10%, 15% and 20% margins.")
heading("Allocation algorithm", 2)
code(["For each eligible forecast origin and chosen policy:", "  forecast = saved forecast for the next interval", "  capacity_mbps = max(0, forecast) * (1 + margin)", "  deficit_mbps = max(actual_mbps - capacity_mbps, 0)", "  excess_mbps = max(capacity_mbps - actual_mbps, 0)", "  served_mbps = min(actual_mbps, capacity_mbps)", "  deficit_mb = deficit_mbps * 600 / 8"])
para("Shortfall frequency is the percentage of evaluated intervals with positive deficit. Deficit volume sums interval deficits in decimal MB. Mean capacity measures provisioned resources, while mean over-provisioning measures positive unused capacity. The served-demand-to-capacity ratio is total served demand divided by total capacity over the evaluated intervals. It is an aggregate ratio, not the average of per-interval ratios.")
heading("Validation based margin selection", 2)
para("For each policy, choose the tested margin with the lowest mean validation capacity subject to shortfalls in at most 1% of validation intervals. If no tested margin qualifies, use 20% and retain an explicit failure flag. The selected margin is then applied to test forecasts without using test demand to tune it. A fallback does not satisfy the reliability requirement.")
table(["Hand check", "Demand", "Forecast", "Margin", "Capacity", "Deficit"], [["Flat traffic", "100 Mbps", "100 Mbps", "10%", "110 Mbps", "0 MB"], ["Sudden spike", "200 Mbps", "100 Mbps", "10%", "110 Mbps", "6,750 MB"]], [1.45, 1.1, 1.1, .75, 1.1, 1.35])
para("The spike deficit is 90 Mbps for 600 seconds, giving 6,750 MB. The conversion uses the CESNET interval rather than the 900 seconds associated with the abandoned dataset. Excluded gaps are not filled or integrated into volume totals.")
para("The simulated deficit is an amount of demand that the allocated pool could not serve under this abstraction. It does not measure packet loss. There is no routing model, physical link capacity, queue, controller actuation delay or latency calculation. The experiment studies policy arithmetic and the capacity-shortfall tradeoff under a specified target definition.")

page("Allocation results and capacity tradeoffs")
para("None of the policies meets the 1% validation shortfall target at any tested margin. All default to 20% with the failure flag retained. The following table reports both the validation evidence for that fallback and the corresponding test outcome.")
alloc_data = []
for p in policies:
    a = default_map[p]
    d = ac["selected_defaults"][p]
    alloc_data.append([plabel[p], f"{float(d['validation_shortfall_pct']):.2f}", f"{float(a['shortfall_pct']):.2f}", f"{float(a['mean_capacity_mbps']):.2f}", f"{float(a['deficit_volume_mb']):,.0f}", f"{float(a['mean_overprovision_mbps']):.2f}"])
table(["Policy", "Val shortfall %", "Test shortfall %", "Mean cap Mbps", "Deficit MB", "Unused Mbps"], alloc_data, [1.2, 1.1, 1.1, 1.05, 1.35, 1.05])
para("Table 4. Allocation at the 20% fallback margin. All test policies use 5,554 identical intervals. Capacity and unused-capacity columns are means; deficit volume is a sum over evaluated intervals.", "Caption")
matrix = []
for p in policies:
    matrix.append([plabel[p]] + [f"{float(next(r for r in allocation if r['policy'] == p and r['split'] == 'test' and float(r['margin']) == margin)['shortfall_pct']):.2f}" for margin in (.05, .10, .15, .20)])
table(["Test shortfall %", "5% margin", "10% margin", "15% margin", "20% margin"], matrix, [2.05, 1.2, 1.2, 1.2, 1.2])
para("Table 5. Identical safety margins allow a direct policy comparison. Complete validation and test metrics for every margin are in the 40-row allocation metrics.csv.", "Caption")
para(f"At 20%, LSTM reduces the test shortfall frequency from {float(reactive['shortfall_pct']):.2f}% to {float(lstm['shortfall_pct']):.2f}%, while mean capacity rises from {float(reactive['mean_capacity_mbps']):.2f} to {float(lstm['mean_capacity_mbps']):.2f} Mbps. It also has more unused capacity. Therefore the reduction is not evidence of a cost-free efficiency improvement.")
para("Forecasting rank and allocation rank differ because the policies depend on forecast bias and the upper tail of underprediction, rather than only average squared error. Transformer has the best ten-minute RMSE, but LSTM has the lowest shortfall frequency among the five policies at the common 20% margin. A separate asymmetric loss or calibrated quantile forecast could target shortfall directly; neither was tuned here.")
ratio_text = "; ".join(f"{plabel[p]} {float(default_map[p]['served_demand_to_capacity_ratio']):.3f}" for p in policies)
para("The test served-demand-to-capacity ratios at 20% are: " + ratio_text + ". A higher ratio alone can conceal greater unmet demand, so it must be interpreted with the deficit and shortfall columns.")

page("Software architecture and demonstration")
para("The repository separates data preparation, model definitions, experiments, dashboard code and tests. Commands run as Python modules from the project root. Data preparation creates common windows, model commands produce saved forecasts, allocation reads those forecasts, and the dashboard reads validated result files. This separation prevents a routine dashboard interaction from silently retraining models or changing the experiment.")
table(["Component", "Input", "Output"], [["Data processing", "Raw institution and time mapping", "Flags, windows, splits and scalers"], ["Forecast experiments", "Prepared windows and configuration", "Checkpoints, predictions and metrics"], ["Allocation", "Aligned ten-minute forecasts", "Margins, replay, metrics and defaults"], ["Dashboard", "Saved CSV and JSON artifacts", "Plots, controls, replay and downloads"], ["Verification", "Saved inputs and result signatures", "Consistency and checkpoint checks"]], [1.55, 2.5, 2.8])
para("The local Streamlit interface offers model, forecast horizon and validation/test selection. Forecast charts use real UTC target times and preserve excluded gaps. The chart-history and replay controls reveal historical intervals without changing full-split metric cards. The comparison view reports forecasting accuracy and timing; CSV downloads allow independent inspection.")
para("The allocation view has its own policy and safety-margin controls. It shows capacity, observed demand and deficits, while preserving the validation-selected fallback warning. Thirty-minute and sixty-minute forecasting views do not imply that those allocation horizons have been evaluated. Allocation stays explicitly ten-minute.")
para("The dashboard checks target identities, observed values, horizon leads, common preparation signatures and recomputed forecast scores. Allocation also checks saved source hashes, metrics and defaults. Missing or inconsistent artifacts produce an explanatory message. Viewing saved forecasts does not import PyTorch or require a GPU.")
heading("Demonstration sequence", 2)
para("Start with the ten-minute test comparison and show that the lowest RMSE and lowest MAPE need not identify the same model. Change to sixty minutes to discuss the GRU result and Transformer degradation. Return to the ten-minute allocation view, compare equal margins and replay a high-demand interval. Explain the extra-capacity tradeoff and the unmet 1% target before presenting limitations.")
code(["source VirtEnv/Scripts/activate", "python -m streamlit run dashboard/app.py", "# Open http://127.0.0.1:8501 in the local browser"])

page("Reproducibility and artifact traceability")
para("The tested training environment records Python 3.12.10, NumPy 2.3.0, Matplotlib 3.10.0 and PyTorch 2.14.1+cu130 on Windows with an RTX 4050 Laptop GPU. Dashboard dependencies are Streamlit 1.66.0, pandas 3.0.6 and Plotly 7.1.0. Dependency files separate preparation, training, CUDA and dashboard installation. No TensorFlow dependency is required.")
heading("Documented workflow", 2)
para("Activate the environment from Git Bash and install the appropriate training and dashboard requirements. The reproduction command obtains data when needed, prepares windows, evaluates both baselines and all neural models, and runs allocation in dependency order. Fresh regeneration backs up existing generated outputs so earlier results remain recoverable. Exact numerical identity on another machine is not promised.")
code(["source VirtEnv/Scripts/activate", "python -m pip install -r requirements-cuda.txt", "python -m pip install -r requirements-dashboard.txt", "python -m scripts.reproduce --download --fresh --device auto", "python -m scripts.verify_results --checkpoints --device auto", "python -m unittest discover -s tests -v", "python -m pip check"])
para("For CPU execution, install requirements-training.txt instead of the CUDA-specific file and use --device cpu. Strict checkpoint replay should use the training backend: replaying saved CUDA checkpoints on CPU introduces numerical differences beyond the strict tolerance. A CPU-only user can verify artifact arithmetic and provenance without --checkpoints, or regenerate the experiment on CPU. A fresh clone omits ignored data and results; the submission package supplies the experiment artifacts separately.")
heading("Result sources", 2)
table(["Reported evidence", "Authoritative local artifact"], [["Target counts and preprocessing", "data/prepared/window_counts.csv and preparation_config.json"], ["All forecast scores and timings", "results/<family>/<horizon>min/metrics.csv"], ["Checkpoint and training details", "Corresponding run_config.json and training_history.csv"], ["Forecast identities and values", "Corresponding predictions.csv"], ["All allocation margins", "results/allocation/10min/metrics.csv"], ["Default test allocation", "results/allocation/10min/default_test_metrics.csv"], ["Report provenance", "deliverables/report/source_manifest.json"]], [2.4, 4.45])
para("The report generator loads metrics from the model result files, enforces complete forecasting and allocation tables, and copies their full-precision values into accompanying CSVs. Its source manifest records SHA-256 hashes of every consumed artifact. Thus rounded report values can be traced to exact local inputs. Rebuilding the report changes its numerical tables when source results change; model retraining is a separate explicit operation.")

page("Limitations and conclusions")
heading("Limits on the findings", 2)
para("One institution cannot represent all ISP, cloud or data-center traffic. The institution choice was informed by availability and prior inspection, but performance on other institutions remains unknown. One fixed seed supplies no estimate of initialization uncertainty. The experiment does not include repeated splits, multiple seeds, formal significance testing or forecast confidence intervals. Small performance gaps should therefore be interpreted cautiously.")
para("Measurement uncertainty remains despite conservative exclusions. Timestamp discontinuities have been guarded rather than fully explained. The buffered outage policy may remove useful observations and cannot guarantee that retained observations are free from undercapture. Rate conversion assumes nominal 600-second intervals, and live export delay is unknown. No anomaly was deleted solely for being statistically unusual.")
para("The chronological split provides a defensible temporal holdout but captures only one period of distribution change. The raw timeline cut points precede exclusions, leaving fewer usable observations around the monitoring outage. Different horizon target sets slightly complicate cross-horizon comparisons. Input scaling counts training-window values, including observations reused in overlapping windows, consistently with the actual learner inputs.")
para("The model comparison is deliberately compact. It includes neither seasonal persistence nor statistical models such as ARIMA, extensive hyperparameter tuning, probabilistic forecasting or covariate ablation. Training and inference times come from individual runs and exclude full-system latency. These omissions bound the conclusions rather than invalidate the implemented comparisons.")
para("The allocation simulation has unlimited abstract provisionable capacity and no switching penalty, queue, routing constraint, pricing model or service-level enforcement. Deficit volume is a simulated unserved amount, not measured packet loss. Equal percentage margins support a transparent comparison but do not satisfy the stated 1% validation reliability target.")
heading("Conclusions", 2)
para("The project demonstrates a complete traffic-to-forecast-to-capacity pipeline with traceable experimental artifacts and a working local interface. Neural models improve test RMSE over persistence at all three leads. Transformer performs best on ten-minute RMSE, while GRU performs best on thirty-minute and sixty-minute RMSE. LSTM remains useful but does not have a universal advantage.")
para("Forecast-based allocation reduces shortfall frequency in the examined test period, while consuming additional capacity. The failed validation reliability target is a central finding: average forecast accuracy alone is insufficient to guarantee dependable provisioning. Future evaluation should add institutions, repeat seeds, test probabilistic or quantile forecasts, and assess policies under realistic capacity, cost and actuation constraints before a live deployment claim.")

page("References and result appendix")
refs = [
    ("[1] Josef Koumar, Karel Hynek, Tomáš Čejka and Pavel Šiška. CESNET-TimeSeries24: Time Series Dataset for Network Traffic Anomaly Detection and Forecasting. Version v1. Zenodo, 2024. Dataset DOI 10.5281/zenodo.13382427. Licensed CC BY 4.0.", "https://zenodo.org/records/13382427"),
    ("[2] Josef Koumar, Karel Hynek, Tomáš Čejka and Pavel Šiška. CESNET-TimeSeries24: Time Series Dataset for Network Traffic Anomaly Detection and Forecasting. Scientific Data 12, article 338, 2025. DOI 10.1038/s41597-025-04603-x. Collection methodology also reviewed in arXiv:2409.18874v1, 2024.", "https://arxiv.org/html/2409.18874v1"),
    ("[3] Sepp Hochreiter and Jürgen Schmidhuber. Long Short-Term Memory. Neural Computation 9(8), 1735-1780, 1997. DOI 10.1162/neco.1997.9.8.1735.", "https://people.idsia.ch/~juergen/lstm1997-2024head.pdf"),
    ("[4] Kyunghyun Cho, Bart van Merrienboer, Caglar Gulcehre, Dzmitry Bahdanau, Fethi Bougares, Holger Schwenk and Yoshua Bengio. Learning Phrase Representations using RNN Encoder-Decoder for Statistical Machine Translation. EMNLP, 2014. arXiv:1406.1078.", "https://arxiv.org/abs/1406.1078"),
    ("[5] Ashish Vaswani, Noam Shazeer, Niki Parmar, Jakob Uszkoreit, Llion Jones, Aidan N. Gomez, Lukasz Kaiser and Illia Polosukhin. Attention Is All You Need. Advances in Neural Information Processing Systems, 2017. arXiv:1706.03762.", "https://arxiv.org/abs/1706.03762"),
]
for desc, url in refs:
    para(desc)
    p = para(url)
    for run in p.runs:
        run.font.size = Pt(9)
heading("Machine readable appendix", 2)
para("forecast_metrics.csv accompanies this report with all 30 model, horizon and split combinations, original metric precision, MAPE inclusion counts and timings. allocation_metrics.csv contains all 40 policy, margin and split combinations. allocation_defaults.csv records all five default test policies. These files preserve details that are rounded or summarized in the printed tables.")
para("Source citations [1] and [2] document the dataset and its collection. Sources [3] to [5] establish the model families. They do not substantiate the numerical results of this implementation; those results are supported by the local artifact paths and hashes listed in source_manifest.json. Dataset attribution is retained when distributing the extracted sample and derived tables.")

doc.save(OUT / "project_report.docx")
(DOCS / "project_report.md").write_text("\n".join(md), encoding="utf-8")
metric_list = [metrics[h, m, split] for h in (10, 30, 60) for m in MODELS for split in ("validation", "test")]
for name, data in (("forecast_metrics.csv", metric_list), ("allocation_metrics.csv", allocation), ("allocation_defaults.csv", defaults)):
    with (OUT / name).open("w", newline="", encoding="utf-8") as f:
        writer = csv.DictWriter(f, fieldnames=list(data[0]))
        writer.writeheader()
        writer.writerows(data)
(OUT / "source_manifest.json").write_text(json.dumps({"report_date": "2026-10-22", "source_sha256": sources, "forecast_rows": len(metric_list), "allocation_rows": len(allocation), "default_rows": len(defaults)}, indent=2) + "\n", encoding="utf-8")
print(f"Created {OUT / 'project_report.docx'}")
print(f"Created {DOCS / 'project_report.md'}")
